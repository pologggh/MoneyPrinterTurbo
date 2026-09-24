from __future__ import annotations

import io
import re
from typing import Any, Sequence
import unicodedata

from bs4 import BeautifulSoup
import docx
from pydantic import BaseModel, ConfigDict, Field
import pdfplumber

from app.domain.evidence import (
    DocumentTextNotExtractableError,
    EvidenceLocator,
    SourceContentUnavailableError,
)

DEFAULT_PROCESSING_VERSION = "knowledge_processing_v1"


def normalize_text(text: str) -> str:
    """Deterministic normalization: NFKC unicode, CRLF to LF, and whitespace trimming.

    Preserves exact factual source wording without summarization or paraphrasing.
    """
    if not text:
        return ""
    # Unicode NFKC normalization
    norm = unicodedata.normalize("NFKC", text)
    # Convert CRLF and CR to LF
    norm = norm.replace("\r\n", "\n").replace("\r", "\n")
    # Clean lines and horizontal whitespace
    lines = norm.split("\n")
    result: list[str] = []
    blank_count = 0
    for line in lines:
        cleaned_line = re.sub(r"[^\S\n]+", " ", line).strip()
        if not cleaned_line:
            blank_count += 1
            if blank_count <= 1 and result:
                result.append("")
        else:
            blank_count = 0
            result.append(cleaned_line)
    return "\n".join(result).strip()


class StructuralBlock(BaseModel):
    """An atomic structural unit of parsed content (heading, paragraph, page, etc.)."""
    model_config = ConfigDict(frozen=True)

    block_index: int
    block_type: str  # HEADING, PARAGRAPH, CODE, LIST_ITEM, PAGE
    text: str
    locator: EvidenceLocator = Field(default_factory=EvidenceLocator)


class ParsedKnowledgeDocument(BaseModel):
    """Normalized structured representation of a source document."""
    model_config = ConfigDict(frozen=True)

    source_document_id: str
    processing_version: str = DEFAULT_PROCESSING_VERSION
    title: str | None = None
    blocks: tuple[StructuralBlock, ...] = Field(default_factory=tuple)
    parser_metadata: dict[str, Any] = Field(default_factory=dict)


class DocumentParser:
    """Unified document parser handling TEXT, Markdown, PDF, DOCX, and HTML."""

    def __init__(self, processing_version: str = DEFAULT_PROCESSING_VERSION) -> None:
        self.processing_version = processing_version

    @staticmethod
    def _resolve_source_and_content(
        arg1: Any,
        arg2: Any = None,
        title: str | None = None,
    ) -> tuple[str, Any, str | None]:
        if hasattr(arg1, "source_document_id"):
            source_id = arg1.source_document_id
            doc_title = title or getattr(arg1, "title", None)
            content = arg2 if arg2 is not None else getattr(arg1, "content_snapshot", "")
            return source_id, content, doc_title
        if hasattr(arg2, "source_document_id"):
            source_id = arg2.source_document_id
            doc_title = title or getattr(arg2, "title", None)
            content = arg1
            return source_id, content, doc_title
        if isinstance(arg2, (bytes, bytearray)):
            return str(arg1), arg2, title
        if isinstance(arg1, (bytes, bytearray)):
            return str(arg2) if arg2 is not None else "src_unknown", arg1, title
        if arg2 is None:
            return "src_unknown", arg1, title
        s1 = str(arg1)
        s2 = str(arg2)
        if "\n" in s1 or len(s1) > len(s2):
            return s2, s1, title
        return s1, s2, title

    def parse_text(
        self,
        arg1: Any,
        arg2: Any = None,
        title: str | None = None,
    ) -> ParsedKnowledgeDocument:
        """Parses raw plain text into paragraph-level structural blocks."""
        source_document_id, text, doc_title = self._resolve_source_and_content(arg1, arg2, title)
        cleaned = normalize_text(text or "")
        if not cleaned:
            return ParsedKnowledgeDocument(
                source_document_id=source_document_id,
                processing_version=self.processing_version,
                title=doc_title,
                blocks=(),
            )

        # Split on double newlines for paragraphs
        raw_paragraphs = re.split(r"\n\s*\n", cleaned)
        blocks: list[StructuralBlock] = []
        current_line = 1

        for idx, para in enumerate(raw_paragraphs, start=1):
            para_clean = para.strip()
            if not para_clean:
                continue
            line_count = len(para_clean.splitlines())
            locator = EvidenceLocator(
                paragraph=idx,
                line_start=current_line,
                line_end=current_line + line_count - 1,
            )
            blocks.append(
                StructuralBlock(
                    block_index=len(blocks) + 1,
                    block_type="PARAGRAPH",
                    text=para_clean,
                    locator=locator,
                )
            )
            current_line += line_count + 1

        return ParsedKnowledgeDocument(
            source_document_id=source_document_id,
            processing_version=self.processing_version,
            title=title,
            blocks=tuple(blocks),
            parser_metadata={"parser": "TextParser", "paragraph_count": len(blocks)},
        )

    def parse_markdown(
        self,
        arg1: Any,
        arg2: Any = None,
        title: str | None = None,
    ) -> ParsedKnowledgeDocument:
        """Parses Markdown text into structural blocks preserving heading hierarchy."""
        source_document_id, markdown_text, doc_title = self._resolve_source_and_content(arg1, arg2, title)
        cleaned = normalize_text(markdown_text or "")
        if not cleaned:
            return ParsedKnowledgeDocument(
                source_document_id=source_document_id,
                processing_version=self.processing_version,
                title=doc_title,
                blocks=(),
            )

        lines = cleaned.split("\n")
        blocks: list[StructuralBlock] = []
        current_section_stack: list[str] = []

        i = 0
        p_index = 0
        while i < len(lines):
            line = lines[i]
            # Check for Markdown headings: # H1, ## H2, etc.
            heading_match = re.match(r"^(#{1,6})\s+(.+)$", line)
            if heading_match:
                level = len(heading_match.group(1))
                h_text = heading_match.group(2).strip()
                if level == 1 and not doc_title:
                    doc_title = h_text

                # Adjust section stack according to heading depth
                if len(current_section_stack) >= level:
                    current_section_stack = current_section_stack[: level - 1]
                current_section_stack.append(h_text)
                section_path = " > ".join(current_section_stack)

                blocks.append(
                    StructuralBlock(
                        block_index=len(blocks) + 1,
                        block_type="HEADING",
                        text=h_text,
                        locator=EvidenceLocator(
                            section=section_path,
                            line_start=i + 1,
                            line_end=i + 1,
                        ),
                    )
                )
                i += 1
                continue

            # Code block
            if line.startswith("```"):
                code_lines = [line]
                start_line = i + 1
                i += 1
                while i < len(lines) and not lines[i].startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                if i < len(lines):
                    code_lines.append(lines[i])
                    i += 1
                code_text = "\n".join(code_lines)
                section_path = " > ".join(current_section_stack) if current_section_stack else None
                blocks.append(
                    StructuralBlock(
                        block_index=len(blocks) + 1,
                        block_type="CODE",
                        text=code_text,
                        locator=EvidenceLocator(
                            section=section_path,
                            line_start=start_line,
                            line_end=i,
                        ),
                    )
                )
                continue

            # Paragraph or list item
            if line.strip():
                para_lines = [line]
                start_line = i + 1
                i += 1
                while i < len(lines) and lines[i].strip() and not lines[i].startswith("#") and not lines[i].startswith("```"):
                    para_lines.append(lines[i])
                    i += 1
                para_text = "\n".join(para_lines).strip()
                p_index += 1
                section_path = " > ".join(current_section_stack) if current_section_stack else None
                b_type = "LIST_ITEM" if para_text.startswith(("- ", "* ", "1. ")) else "PARAGRAPH"
                blocks.append(
                    StructuralBlock(
                        block_index=len(blocks) + 1,
                        block_type=b_type,
                        text=para_text,
                        locator=EvidenceLocator(
                            section=section_path,
                            paragraph=p_index,
                            line_start=start_line,
                            line_end=i,
                        ),
                    )
                )
            else:
                i += 1

        return ParsedKnowledgeDocument(
            source_document_id=source_document_id,
            processing_version=self.processing_version,
            title=doc_title,
            blocks=tuple(blocks),
            parser_metadata={"parser": "MarkdownParser", "block_count": len(blocks)},
        )

    def parse_pdf(
        self,
        arg1: Any,
        arg2: Any = None,
        title: str | None = None,
    ) -> ParsedKnowledgeDocument:
        """Parses a PDF file page by page using pdfplumber.

        Raises DocumentTextNotExtractableError if the document is scanned or image-only.
        """
        source_document_id, pdf_bytes, doc_title = self._resolve_source_and_content(arg1, arg2, title)
        if not pdf_bytes:
            raise SourceContentUnavailableError(f"PDF content for '{source_document_id}' is empty.")

        blocks: list[StructuralBlock] = []
        total_extracted_chars = 0

        try:
            with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                for page_idx, page in enumerate(pdf.pages, start=1):
                    page_text = page.extract_text() or ""
                    cleaned_page = normalize_text(page_text)
                    if cleaned_page:
                        total_extracted_chars += len(cleaned_page)
                        # Split into paragraphs per page
                        page_paras = re.split(r"\n\s*\n", cleaned_page)
                        for p_idx, para in enumerate(page_paras, start=1):
                            clean_para = para.strip()
                            if clean_para:
                                blocks.append(
                                    StructuralBlock(
                                        block_index=len(blocks) + 1,
                                        block_type="PARAGRAPH",
                                        text=clean_para,
                                        locator=EvidenceLocator(
                                            page=page_idx,
                                            paragraph=p_idx,
                                        ),
                                    )
                                )
        except Exception as exc:
            if isinstance(exc, DocumentTextNotExtractableError):
                raise
            raise SourceContentUnavailableError(f"Failed to read PDF '{source_document_id}': {exc}") from exc

        if total_extracted_chars == 0 or not blocks:
            raise DocumentTextNotExtractableError(
                f"DOCUMENT_TEXT_NOT_EXTRACTABLE: PDF document '{source_document_id}' contains no "
                f"extractable digital text (scanned or image-only PDF)."
            )

        return ParsedKnowledgeDocument(
            source_document_id=source_document_id,
            processing_version=self.processing_version,
            title=doc_title,
            blocks=tuple(blocks),
            parser_metadata={"parser": "PdfParser", "extracted_chars": total_extracted_chars, "page_count": len(pdf.pages)},
        )

    def parse_docx(
        self,
        arg1: Any,
        arg2: Any = None,
        title: str | None = None,
    ) -> ParsedKnowledgeDocument:
        """Parses a DOCX document into paragraphs and headings preserving structure."""
        source_document_id, docx_bytes, doc_title = self._resolve_source_and_content(arg1, arg2, title)
        if not docx_bytes:
            raise SourceContentUnavailableError(f"DOCX content for '{source_document_id}' is empty.")

        try:
            doc = docx.Document(io.BytesIO(docx_bytes))
        except Exception as exc:
            raise SourceContentUnavailableError(f"Failed to read DOCX '{source_document_id}': {exc}") from exc

        blocks: list[StructuralBlock] = []
        current_section_stack: list[str] = []
        p_index = 0

        for p in doc.paragraphs:
            text = normalize_text(p.text)
            if not text:
                continue

            style_name = (p.style.name or "").lower() if p.style else ""
            if "heading" in style_name:
                level_match = re.search(r"\d+", style_name)
                level = int(level_match.group(0)) if level_match else 1
                if not doc_title and level == 1:
                    doc_title = text

                if len(current_section_stack) >= level:
                    current_section_stack = current_section_stack[: level - 1]
                current_section_stack.append(text)
                section_path = " > ".join(current_section_stack)

                blocks.append(
                    StructuralBlock(
                        block_index=len(blocks) + 1,
                        block_type="HEADING",
                        text=text,
                        locator=EvidenceLocator(section=section_path),
                    )
                )
            else:
                p_index += 1
                section_path = " > ".join(current_section_stack) if current_section_stack else None
                blocks.append(
                    StructuralBlock(
                        block_index=len(blocks) + 1,
                        block_type="PARAGRAPH",
                        text=text,
                        locator=EvidenceLocator(
                            section=section_path,
                            paragraph=p_index,
                        ),
                    )
                )

        return ParsedKnowledgeDocument(
            source_document_id=source_document_id,
            processing_version=self.processing_version,
            title=doc_title,
            blocks=tuple(blocks),
            parser_metadata={"parser": "DocxParser", "paragraph_count": len(blocks)},
        )

    def parse_html(
        self,
        arg1: Any,
        arg2: Any = None,
        title: str | None = None,
    ) -> ParsedKnowledgeDocument:
        """Parses HTML content into structural blocks (headings, paragraphs)."""
        source_document_id, html_text, doc_title = self._resolve_source_and_content(arg1, arg2, title)
        if not html_text:
            return ParsedKnowledgeDocument(
                source_document_id=source_document_id,
                processing_version=self.processing_version,
                title=doc_title,
                blocks=(),
            )
        soup = BeautifulSoup(html_text, "html.parser")
        if title is None and soup.title and soup.title.string:
            extracted_title = soup.title.string.strip()
            if extracted_title:
                doc_title = extracted_title
        elif not doc_title and soup.title and soup.title.string:
            doc_title = soup.title.string.strip()

        for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript", "svg"]):
            tag.decompose()

        blocks: list[StructuralBlock] = []
        current_section_stack: list[str] = []
        p_index = 0

        for elem in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li"]):
            text = normalize_text(elem.get_text())
            if not text:
                continue

            tag_name = elem.name.lower()
            if tag_name.startswith("h"):
                level = int(tag_name[1])
                if not doc_title and level == 1:
                    doc_title = text

                if len(current_section_stack) >= level:
                    current_section_stack = current_section_stack[: level - 1]
                current_section_stack.append(text)
                section_path = " > ".join(current_section_stack)

                blocks.append(
                    StructuralBlock(
                        block_index=len(blocks) + 1,
                        block_type="HEADING",
                        text=text,
                        locator=EvidenceLocator(section=section_path),
                    )
                )
            else:
                p_index += 1
                section_path = " > ".join(current_section_stack) if current_section_stack else None
                b_type = "LIST_ITEM" if tag_name == "li" else "PARAGRAPH"
                blocks.append(
                    StructuralBlock(
                        block_index=len(blocks) + 1,
                        block_type=b_type,
                        text=text,
                        locator=EvidenceLocator(section=section_path, paragraph=p_index),
                    )
                )

        return ParsedKnowledgeDocument(
            source_document_id=source_document_id,
            processing_version=self.processing_version,
            title=doc_title,
            blocks=tuple(blocks),
            parser_metadata={"parser": "HtmlParser", "block_count": len(blocks)},
        )

    def parse_source_document(
        self,
        source_document: Any,
        raw_content: bytes | str | None = None,
    ) -> ParsedKnowledgeDocument:
        """Dispatches source document parsing to the correct parser based on source type and media type."""
        payload = raw_content if raw_content is not None else getattr(source_document, "content_snapshot", None)
        media_type = (getattr(source_document, "media_type", None) or "").lower()
        source_type = str(getattr(source_document, "source_type", ""))
        locator = (getattr(source_document, "source_locator", None) or "").lower()

        if payload is None:
            raise SourceContentUnavailableError(
                f"Source document '{source_document.source_document_id}' has no content snapshot and no raw content was provided."
            )

        if "pdf" in media_type or locator.endswith(".pdf"):
            bytes_data = payload.encode("utf-8") if isinstance(payload, str) else payload
            return self.parse_pdf(source_document, bytes_data)
        elif "word" in media_type or "docx" in media_type or locator.endswith(".docx"):
            bytes_data = payload.encode("utf-8") if isinstance(payload, str) else payload
            return self.parse_docx(source_document, bytes_data)
        elif "markdown" in media_type or locator.endswith(".md") or locator.endswith(".markdown"):
            str_data = payload if isinstance(payload, str) else payload.decode("utf-8", errors="replace")
            return self.parse_markdown(source_document, str_data)
        elif "html" in media_type or locator.endswith(".html") or locator.endswith(".htm") or source_type == "URL":
            str_data = payload if isinstance(payload, str) else payload.decode("utf-8", errors="replace")
            return self.parse_html(source_document, str_data)
        else:
            str_data = payload if isinstance(payload, str) else payload.decode("utf-8", errors="replace")
            return self.parse_text(source_document, str_data)
