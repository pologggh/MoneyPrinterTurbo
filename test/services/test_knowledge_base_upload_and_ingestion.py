from __future__ import annotations

import os
from pathlib import Path
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.application.knowledge_base_service import (
    FileSizeExceededError,
    InvalidFileFormatError,
    KnowledgeBaseCommandService,
    SecurityValidationError,
    sanitize_filename,
)
from app.application.knowledge_retrieval_service import KnowledgeProcessingService
from app.domain.evidence import SourceStatus
from app.persistence.models import Base
from app.services.knowledge.embedding_provider import DeterministicFakeEmbeddingProvider


import tempfile


@pytest.fixture
def tmp_path():
    with tempfile.TemporaryDirectory() as td:
        yield Path(td)


@pytest.fixture
def session_factory():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    return SessionLocal


def test_sanitize_filename_security():
    # Path traversal attempts
    assert sanitize_filename("../../etc/passwd.txt") == "passwd.txt"
    assert sanitize_filename("..\\..\\windows\\system32\\cmd.exe") == "cmd.exe"
    assert sanitize_filename("normal_file.pdf") == "normal_file.pdf"
    assert sanitize_filename("my:file?name*.docx") == "my_file_name_.docx"


def test_file_format_and_size_validation(tmp_path, session_factory, monkeypatch):
    with session_factory() as session:
        proc_service = KnowledgeProcessingService(
            session=session,
            embedding_provider=DeterministicFakeEmbeddingProvider(dimension=16),
        )
        service = KnowledgeBaseCommandService(
            session=session,
            storage_dir=tmp_path,
            processing_service=proc_service,
        )

        kb = service.create_knowledge_base("Security Test KB")

        # 1. Unsupported extension
        with pytest.raises(InvalidFileFormatError, match="Unsupported file format"):
            service.upload_document(
                kb_id=kb.knowledge_base_id,
                file_bytes=b"malicious content",
                filename="script.sh",
            )

        # 2. File size exceeded (tested with 1KB monkeypatch to avoid 50MB buffer allocation)
        monkeypatch.setattr("app.application.knowledge_base_service.MAX_FILE_SIZE_BYTES", 1024)
        huge_bytes = b"0" * 1025
        with pytest.raises(FileSizeExceededError, match="exceeds maximum limit"):
            service.upload_document(
                kb_id=kb.knowledge_base_id,
                file_bytes=huge_bytes,
                filename="huge_file.txt",
            )


def test_upload_and_parse_txt_md_html(tmp_path, session_factory):
    with session_factory() as session:
        proc_service = KnowledgeProcessingService(
            session=session,
            embedding_provider=DeterministicFakeEmbeddingProvider(dimension=16),
        )
        service = KnowledgeBaseCommandService(
            session=session,
            storage_dir=tmp_path,
            processing_service=proc_service,
        )

        kb = service.create_knowledge_base("Multi-format KB")

        # 1. TXT
        txt_content = b"Artificial Intelligence and Machine Learning are transforming modern science."
        doc_txt = service.upload_document(
            kb_id=kb.knowledge_base_id,
            file_bytes=txt_content,
            filename="intro.txt",
        )
        assert doc_txt.status == SourceStatus.READY

        # 2. Markdown
        md_content = b"# Deep Learning\n\nNeural networks learn hierarchical representations from data."
        doc_md = service.upload_document(
            kb_id=kb.knowledge_base_id,
            file_bytes=md_content,
            filename="notes.md",
        )
        assert doc_md.status == SourceStatus.READY

        # 3. HTML
        html_content = b"<html><body><h1>Quantum Computing</h1><p>Qubits utilize superposition.</p></body></html>"
        doc_html = service.upload_document(
            kb_id=kb.knowledge_base_id,
            file_bytes=html_content,
            filename="quantum.html",
        )
        assert doc_html.status == SourceStatus.READY

        # Verify docs listed
        docs = service.list_documents(kb.knowledge_base_id)
        assert len(docs) == 3
        for d in docs:
            assert d["status"] == "READY"
            assert d["chunk_count"] >= 1


def test_sha256_deduplication(tmp_path, session_factory):
    with session_factory() as session:
        proc_service = KnowledgeProcessingService(
            session=session,
            embedding_provider=DeterministicFakeEmbeddingProvider(dimension=16),
        )
        service = KnowledgeBaseCommandService(
            session=session,
            storage_dir=tmp_path,
            processing_service=proc_service,
        )

        kb = service.create_knowledge_base("Dedup KB")
        content = b"Duplicate file content check."

        doc1 = service.upload_document(kb.knowledge_base_id, content, "file1.txt")
        doc2 = service.upload_document(kb.knowledge_base_id, content, "file2.txt")

        # Must return the identical source document ID
        assert doc1.source_document_id == doc2.source_document_id
        docs = service.list_documents(kb.knowledge_base_id)
        assert len(docs) == 1


def test_upload_docx_and_scanned_pdf_handling(tmp_path, session_factory):
    import io
    import docx
    from unittest.mock import patch
    from app.domain.evidence import DocumentTextNotExtractableError

    with session_factory() as session:
        proc_service = KnowledgeProcessingService(
            session=session,
            embedding_provider=DeterministicFakeEmbeddingProvider(dimension=16),
        )
        service = KnowledgeBaseCommandService(
            session=session,
            storage_dir=tmp_path,
            processing_service=proc_service,
        )

        kb = service.create_knowledge_base("Docx and PDF KB")

        # 1. DOCX upload
        doc_obj = docx.Document()
        doc_obj.add_paragraph("Machine learning systems process data to extract patterns.")
        docx_buf = io.BytesIO()
        doc_obj.save(docx_buf)
        docx_bytes = docx_buf.getvalue()

        doc_docx = service.upload_document(
            kb_id=kb.knowledge_base_id,
            file_bytes=docx_bytes,
            filename="ml_notes.docx",
        )
        assert doc_docx.status == SourceStatus.READY
        chunks = proc_service.evidence_repo.list_chunks_for_source(doc_docx.source_document_id)
        assert len(chunks) >= 1

        # 2. Scanned / image-only PDF (mocking pdfplumber to simulate unextractable PDF)
        with patch.object(proc_service.parser, "parse_pdf", side_effect=DocumentTextNotExtractableError("Scanned PDF with no text")):
            doc_pdf = service.upload_document(
                kb_id=kb.knowledge_base_id,
                file_bytes=b"%PDF-1.4 dummy scanned content",
                filename="scanned.pdf",
            )
            assert doc_pdf.status == SourceStatus.FAILED
            assert "Scanned PDF" in (doc_pdf.metadata_json.get("processing_error") or "")

        # 3. Retry document
        with patch.object(proc_service.parser, "parse_pdf") as mock_parse:
            from app.domain.evidence import EvidenceLocator
            from app.services.knowledge.document_parser import ParsedKnowledgeDocument, StructuralBlock
            mock_parse.return_value = ParsedKnowledgeDocument(
                source_document_id=doc_pdf.source_document_id,
                title="OCR Fixed",
                blocks=(StructuralBlock(block_index=1, block_type="PARAGRAPH", text="Extracted OCR text successfully."),),
            )
            retried_chunks = service.retry_document(kb.knowledge_base_id, doc_pdf.source_document_id)
            assert len(retried_chunks) >= 1
            retried_doc = proc_service.evidence_repo.get_source_document(doc_pdf.source_document_id)
            assert retried_doc.status == SourceStatus.READY
